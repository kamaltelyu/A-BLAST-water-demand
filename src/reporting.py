from __future__ import annotations

import json
from pathlib import Path
from typing import Dict

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def fmt_pm(mean, sd, digits=6):
    if pd.isna(sd):
        return f"{mean:.{digits}f}"
    return f"{mean:.{digits}f} ± {sd:.{digits}f}"


def _set_doc_font(doc: Document, font_name="Arial", size=9):
    for sty_name in ["Normal", "Table Normal"]:
        if sty_name in doc.styles:
            doc.styles[sty_name].font.name = font_name
            doc.styles[sty_name].font.size = Pt(size)


def _save_df_docx(df: pd.DataFrame, title: str, path: Path, note: str | None = None):
    doc = Document()
    _set_doc_font(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    for i, c in enumerate(df.columns):
        table.rows[0].cells[i].text = str(c)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, v in enumerate(row.tolist()):
            cells[i].text = str(v)
    if note:
        p2 = doc.add_paragraph(note)
        if p2.runs:
            p2.runs[0].italic = True
    doc.save(path)


def make_table1(metrics_summary: pd.DataFrame, out_csv: Path, out_docx: Path):
    rows = []
    order = ["BLS", "Transformer", "A-BLAST"]
    for ds in ["Chennai", "Melbourne"]:
        sub = metrics_summary[metrics_summary.dataset.str.lower() == ds.lower()]
        for H in sorted(sub.horizon.unique()):
            for model in order:
                r = sub[(sub.horizon == H) & (sub.model == model)]
                if r.empty:
                    continue
                r = r.iloc[0]
                rows.append({
                    "Dataset": ds,
                    "Horizon": int(H),
                    "Model": model,
                    "MSE (mean ± SD)": fmt_pm(r.mse_mean, r.mse_sd),
                    "MAE (mean ± SD)": fmt_pm(r.mae_mean, r.mae_sd),
                    "RMSE (mean ± SD)": fmt_pm(r.rmse_mean, r.rmse_sd),
                })
    df = pd.DataFrame(rows)
    df.to_csv(out_csv, index=False)
    _save_df_docx(
        df,
        "Table 1. Rolling-origin forecasting performance across datasets and horizons.",
        out_docx,
        "For each of five pre-specified test origins, metrics were first averaged across the five locked random seeds. Values shown are mean ± SD across the five temporal origins; lower values indicate better performance.",
    )
    return df


def make_table2(timing_summary: pd.DataFrame, out_csv: Path, out_docx: Path):
    rows = []
    for ds in ["Chennai", "Melbourne"]:
        sub = timing_summary[timing_summary.dataset.str.lower() == ds.lower()]
        for H in sorted(sub.horizon.unique()):
            trow = sub[(sub.horizon == H) & (sub.model == "Transformer")]
            if trow.empty:
                continue
            tref = float(trow.iloc[0].training_time_mean_s)
            for model in ["BLS", "Transformer", "A-BLAST"]:
                r = sub[(sub.horizon == H) & (sub.model == model)]
                if r.empty:
                    continue
                r = r.iloc[0]
                speed = tref / float(r.training_time_mean_s) if float(r.training_time_mean_s) > 0 else np.nan
                rows.append({
                    "Dataset": ds,
                    "Horizon": int(H),
                    "Model": model,
                    "Train time, s (mean ± SD)": fmt_pm(r.training_time_mean_s, r.training_time_sd_s, 4),
                    "Inference time/origin, s (mean ± SD)": fmt_pm(r.prediction_time_mean_s, r.prediction_time_sd_s, 6),
                    "Train + mean inference, s": fmt_pm(r.total_time_mean_s, r.total_time_sd_s, 4),
                    "Training speed-up vs Transformer": f"{speed:.2f}×" if np.isfinite(speed) else "NA",
                })
    df = pd.DataFrame(rows)
    df.to_csv(out_csv, index=False)
    _save_df_docx(
        df,
        "Table 2. Computational efficiency across datasets and horizons.",
        out_docx,
        "Training time is measured once per fitted seed-level model and summarized across five seeds. Inference time is the seed-level mean of median model-inference times across five rolling test origins after warm-up. Common STL preprocessing is excluded from model timing. Training speed-up is Transformer mean training time divided by the corresponding model mean training time.",
    )
    return df


def make_table3(metrics_by_seed_origin: pd.DataFrame, out_csv: Path, out_docx: Path):
    rows = []
    keys = ["dataset", "horizon", "origin_id", "seed"]
    for baseline in ["BLS", "Transformer"]:
        for metric in ["mae", "mse", "rmse"]:
            a = metrics_by_seed_origin[metrics_by_seed_origin.model == "A-BLAST"][keys + [metric]]
            b = metrics_by_seed_origin[metrics_by_seed_origin.model == baseline][keys + [metric]]
            m = a.merge(b, on=keys, suffixes=("_ablast", "_base"))
            base = m[f"{metric}_base"].to_numpy(dtype=float)
            ab = m[f"{metric}_ablast"].to_numpy(dtype=float)
            pct = 100.0 * (ab - base) / base
            diff = ab - base
            rows.append({
                "Comparison": f"A-BLAST vs {baseline}",
                "Metric": metric.upper(),
                "Mean paired change (%)": f"{np.mean(pct):.2f}%",
                "Median paired change (%)": f"{np.median(pct):.2f}%",
                "A-BLAST wins": int((diff < 0).sum()),
                "Ties": int(np.isclose(diff, 0.0).sum()),
                "Total matched evaluations": int(len(m)),
            })
    df = pd.DataFrame(rows)
    df.to_csv(out_csv, index=False)
    _save_df_docx(
        df,
        "Table 3. Unit-free paired robustness comparison across rolling origins.",
        out_docx,
        "Paired percentage change is 100 × (A-BLAST error − baseline error) / baseline error; negative values favor A-BLAST. Pairs are matched by dataset, horizon, rolling origin, and random seed. Percentage changes are unit-free, avoiding invalid aggregation of Chennai (mcft) and Melbourne (ML) errors. The table is descriptive; no independence-based p-value is claimed.",
    )
    return df


def make_table4(ablation_summary: pd.DataFrame, out_csv: Path, out_docx: Path):
    rows = []
    variant_order = ["BLS baseline", "w/o attention", "w/o Transformer", "w/o enhancement", "Full A-BLAST"]
    for ds in ["Chennai", "Melbourne"]:
        sub = ablation_summary[ablation_summary.dataset.str.lower() == ds.lower()]
        for variant in variant_order:
            r = sub[sub.variant == variant]
            if r.empty:
                continue
            r = r.iloc[0]
            rows.append({
                "Dataset": ds,
                "Horizon": int(r.horizon),
                "Variant": variant,
                "MSE (mean ± SD)": fmt_pm(r.mse_mean, r.mse_sd),
                "MAE (mean ± SD)": fmt_pm(r.mae_mean, r.mae_sd),
                "RMSE (mean ± SD)": fmt_pm(r.rmse_mean, r.rmse_sd),
            })
    df = pd.DataFrame(rows)
    df.to_csv(out_csv, index=False)
    _save_df_docx(
        df,
        "Table 4. Rolling-origin architectural ablation at the locked ablation horizon.",
        out_docx,
        "For each of five test origins, variant metrics were first averaged across five seeds; values are mean ± SD across temporal origins. BLS baseline and Full A-BLAST are reused from the main experiment, while the three removal variants are independently fitted under the same protocol.",
    )
    return df


def _replace_table_value(doc: Document, label: str, value: str):
    found = False
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) >= 2 and row.cells[0].text.strip().lower() == label.strip().lower():
                row.cells[1].text = value
                found = True
    return found


def make_table5_from_template(template_path: Path, out_path: Path, cfg: Dict, env: Dict, selected: Dict):
    doc = Document(template_path)
    values = {
        "Train / validation / test split": f"{int(cfg['train_ratio']*100)}% / {int(cfg['val_ratio']*100)}% / {int((1-cfg['train_ratio']-cfg['val_ratio'])*100)}%, chronological; no temporal shuffling",
        "Input (lookback) window length, L": f"{cfg['lookback']} daily time steps",
        "Forecasting horizons, H": ", ".join(str(x) for x in cfg['horizons']),
        "Normalization": "z-score; scaler fitted only on the fitting history (train for tuning; train+validation for final model fit); future target values never used",
        "Seasonal period": str(cfg['stl_period']),
        "Trend / seasonal smoother": "STL (statsmodels), LOESS-based trend and seasonal smoothers; recomputed using history strictly before each test forecast origin",
        "Robust fitting": "Yes" if cfg['stl_robust'] else "No",
        "Components modeled": "Trend + seasonal; residual excluded from reconstruction model",
        "# BLS feature nodes": str(cfg['feature_nodes']),
        "# enhancement nodes": str(cfg['enhancement_nodes']),
        "Feature-mapping activation": "Sigmoid (σ)",
        "Attention dimension": f"{cfg['feature_nodes']} fixed seeded feature dimensions; {cfg['n_heads']} attention heads",
        "Ridge regularization, λ": f"Selected separately for BLS/A-BLAST and horizon from validation MAE over {list(cfg['ridge_grid'])}; exact selections in selected_hyperparameters.json",
        "# encoder layers": f"A-BLAST fixed block: {cfg['num_layers']}; standalone Transformer: {cfg['transformer_layers']}",
        "# attention heads": f"A-BLAST: {cfg['n_heads']}; standalone Transformer: {cfg['transformer_heads']}",
        "Model / embedding dimension, d": f"A-BLAST: {cfg['feature_nodes']}; standalone Transformer: {cfg['transformer_d_model']}",
        "Feed-forward dimension": f"A-BLAST: {cfg['ff_dim']}; standalone Transformer: {cfg['transformer_ff_dim']}",
        "Dropout": f"A-BLAST fixed block: {cfg['dropout']}; standalone Transformer: {cfg['transformer_dropout']}",
        "Positional encoding": "Sinusoidal",
        "Optimizer": "A-BLAST: none for fixed attention/Transformer block; analytical ridge output. Standalone Transformer: Adam",
        "Learning rate": f"A-BLAST: not applicable; standalone Transformer: {cfg['transformer_lr']}",
        "Batch size": f"A-BLAST: not applicable; standalone Transformer: {cfg['transformer_batch_size']}",
        "Epochs / early stopping": f"A-BLAST: not applicable; standalone Transformer: max {cfg['transformer_max_epochs']} epochs, patience {cfg['transformer_patience']}",
        "# feature / enhancement nodes": f"{cfg['feature_nodes']} / {cfg['enhancement_nodes']}",
        "Search strategy": "Validation-set grid search for ridge λ; all architectural dimensions fixed a priori; test origins excluded from selection",
        "Search ranges": f"ridge λ ∈ {list(cfg['ridge_grid'])}",
        "Random seed(s) / # runs averaged": f"{list(cfg['seeds'])}; five independent fitted models per dataset–horizon–model, each evaluated at {cfg['n_test_origins']} rolling origins ({len(cfg['seeds'])*cfg['n_test_origins']} seed–origin evaluations)",
        "Language & key libraries": f"Python {env.get('python','?')}; PyTorch {env.get('torch','?')}; NumPy {env.get('numpy','?')}; pandas {env.get('pandas','?')}; scikit-learn {env.get('scikit_learn','?')}; statsmodels {env.get('statsmodels','?')}",
        "Hardware (CPU / GPU, RAM)": f"CPU: {env.get('cpu','?')}; GPU: {env.get('gpu','none')} ({env.get('gpu_total_memory_gb','NA')} GB VRAM); RAM: {env.get('ram_total_gb','NA')} GB. Model timing uses time.perf_counter(), CUDA synchronization, {cfg['inference_warmup']} warm-up calls, and {cfg['inference_repeats']} timed inference repeats per origin. Common STL preprocessing is excluded from model timing.",
    }
    for label, val in values.items():
        _replace_table_value(doc, label, val)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "⟨" in cell.text:
                    cell.text = cell.text.replace("⟨fill — required for the efficiency claims⟩", values["Hardware (CPU / GPU, RAM)"])
    # Add an explicit rolling-origin row to the configuration table.
    if doc.tables:
        t = doc.tables[0]
        r = t.add_row().cells
        if len(r) >= 2:
            r[0].text = "Rolling-origin test design"
            r[1].text = (
                f"{cfg['n_test_origins']} evenly spaced feasible forecast origins per horizon within the held-out test interval. "
                "The final model is fitted once on train+validation and remains fixed; earlier test observations may be used only as past observed inputs at later origins."
            )
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("To assess forecasting performance"):
            p.text = (
                f"To assess forecasting performance, A-BLAST was compared with standalone BLS and standalone Transformer across future forecasting horizons "
                f"{', '.join(str(x) for x in cfg['horizons'])}. MSE, MAE, and RMSE were computed on the original target scale. "
                f"A fixed-model rolling-origin backtest with {cfg['n_test_origins']} pre-specified test origins per horizon was used. "
                "All models shared the same chronological split and lookback. Hyperparameters were selected without test data; each final seed-level model was fitted once on train+validation, then evaluated without test-period refitting. "
                "For every origin, STL and the input window were constructed only from observations available before that forecast origin."
            )
        elif txt.startswith("Note: replace each"):
            p.text = "Note: All entries were populated automatically from the locked rerun configuration and captured runtime environment. Exact horizon/model ridge selections and forecast-origin dates are preserved in selected_hyperparameters.json and split_manifest.json."
    doc.save(out_path)


def make_table6_from_template(template_path: Path, out_path: Path, profiles: Dict[str, Dict]):
    doc = Document(template_path)
    table = doc.tables[0]
    mapping = {
        "Source": (profiles["Chennai"]["source_label"], profiles["Melbourne"]["source_label"]),
        "Target variable": (profiles["Chennai"]["target_variable"], profiles["Melbourne"]["target_variable"]),
        "Unit": (profiles["Chennai"]["unit"], profiles["Melbourne"]["unit"]),
        "Temporal coverage (start – end)": (f"{profiles['Chennai']['start']} – {profiles['Chennai']['end']}", f"{profiles['Melbourne']['start']} – {profiles['Melbourne']['end']}"),
        "Sampling frequency": ("Daily", "Daily"),
        "Number of observations": (str(profiles["Chennai"]["n"]), str(profiles["Melbourne"]["n"])),
        "Minimum / Maximum": (f"{profiles['Chennai']['min']:.3f} / {profiles['Chennai']['max']:.3f}", f"{profiles['Melbourne']['min']:.3f} / {profiles['Melbourne']['max']:.3f}"),
        "Mean / Std. dev.": (f"{profiles['Chennai']['mean']:.3f} / {profiles['Chennai']['std']:.3f}", f"{profiles['Melbourne']['mean']:.3f} / {profiles['Melbourne']['std']:.3f}"),
        "Seasonal period (STL)": (str(profiles["Chennai"]["stl_period"]), str(profiles["Melbourne"]["stl_period"])),
        "Missing-value handling": (profiles["Chennai"]["missing_handling"], profiles["Melbourne"]["missing_handling"]),
        "Components modeled": ("Trend + seasonal (residual excluded)", "Trend + seasonal (residual excluded)"),
    }
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label in mapping and len(row.cells) >= 3:
            row.cells[1].text = mapping[label][0]
            row.cells[2].text = mapping[label][1]
    row = table.add_row(); row.cells[0].text = "Missing calendar days"; row.cells[1].text = str(profiles["Chennai"]["missing_calendar_days"]); row.cells[2].text = str(profiles["Melbourne"]["missing_calendar_days"])
    row = table.add_row(); row.cells[0].text = "Raw-file SHA-256"; row.cells[1].text = profiles["Chennai"]["sha256"]; row.cells[2].text = profiles["Melbourne"]["sha256"]
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("The characteristics of the two datasets"):
            p.text = "The characteristics of the two datasets used for validation are summarized in Table 6. All quantitative values are computed automatically from the exact raw files used in the locked rerun."
        elif txt.startswith("Note:"):
            p.text = "Note: Quantitative values and SHA-256 hashes were computed automatically from the exact downloaded raw files used in the rerun. The pipeline aborts if target values or calendar dates are missing; no target imputation is performed."
    doc.save(out_path)


def make_fig2_accuracy_efficiency(metrics: pd.DataFrame, timing: pd.DataFrame, png: Path, pdf: Path):
    fig, axes = plt.subplots(1, 2, figsize=(10.2, 4.4))
    units = {"Chennai": "mcft", "Melbourne": "ML"}
    for ax, ds in zip(axes, ["Chennai", "Melbourne"]):
        m = metrics[metrics.dataset.str.lower() == ds.lower()].groupby("model", as_index=False).agg(mae=("mae", "mean"))
        t = timing[timing.dataset.str.lower() == ds.lower()].groupby("model", as_index=False).agg(total=("total_time_s", "mean"))
        a = m.merge(t, on="model")
        for _, r in a.iterrows():
            ax.scatter(r["total"], r["mae"], s=55)
            ax.annotate(r["model"], (r["total"], r["mae"]), xytext=(6, 5), textcoords="offset points", fontsize=8.5)
        ax.set_xscale("log")
        ax.set_xlabel("Train + mean inference time (s, log scale)")
        ax.set_ylabel(f"MAE ({units[ds]})")
        ax.set_title(ds)
        ax.grid(True, alpha=0.25)
        ax.margins(x=0.16, y=0.12)
    fig.suptitle("Dataset-specific accuracy–efficiency positioning")
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(png, dpi=300, bbox_inches="tight")
    fig.savefig(pdf, bbox_inches="tight")
    plt.close(fig)


def make_fig3_boxplots(metrics: pd.DataFrame, timing: pd.DataFrame, png: Path, pdf: Path):
    order = ["BLS", "Transformer", "A-BLAST"]
    fig, axes = plt.subplots(2, 3, figsize=(11.3, 7.0))
    units = {"Chennai": ("mcft", "mcft²"), "Melbourne": ("ML", "ML²")}
    for i, ds in enumerate(["Chennai", "Melbourne"]):
        dm = metrics[metrics.dataset.str.lower() == ds.lower()]
        dt = timing[timing.dataset.str.lower() == ds.lower()]
        specs = [
            (dm, "mae", f"MAE ({units[ds][0]})", False),
            (dm, "mse", f"MSE ({units[ds][1]})", False),
            (dt, "total_time_s", "Train + mean inference (s)", True),
        ]
        for j, (src, col, label, logy) in enumerate(specs):
            ax = axes[i, j]
            vals = [src.loc[src.model == model, col].to_numpy(dtype=float) for model in order]
            ax.boxplot(vals, tick_labels=order, showfliers=True)
            ax.set_title(label)
            ax.grid(True, axis="y", alpha=0.25)
            ax.tick_params(axis="x", rotation=18)
            if logy:
                ax.set_yscale("log")
            if j == 0:
                ax.set_ylabel(ds)
    fig.suptitle("Dataset-specific error and computational-time distributions")
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(png, dpi=300, bbox_inches="tight")
    fig.savefig(pdf, bbox_inches="tight")
    plt.close(fig)


def make_fig4_forecast(preds: pd.DataFrame, png: Path, pdf: Path, horizon: int = 96, origin_id: str = "O3"):
    models = ["BLS", "Transformer", "A-BLAST"]
    datasets = ["Chennai", "Melbourne"]
    fig, axes = plt.subplots(2, 3, figsize=(12.2, 6.6), sharex=False)
    for i, ds in enumerate(datasets):
        for j, model in enumerate(models):
            ax = axes[i, j]
            sub = preds[(preds.dataset.str.lower() == ds.lower()) & (preds.horizon == horizon) & (preds.model == model) & (preds.origin_id == origin_id)].copy()
            sub["date"] = pd.to_datetime(sub["date"])
            g = sub.groupby("date").agg(y_true=("y_true", "first"), y_mean=("y_pred", "mean"), y_sd=("y_pred", "std")).reset_index()
            ax.plot(g.date, g.y_true, label="Ground truth", linewidth=1.4)
            ax.plot(g.date, g.y_mean, linestyle="--", label="Mean forecast", linewidth=1.2)
            lo = (g.y_mean - g.y_sd.fillna(0)).to_numpy(dtype=float)
            hi = (g.y_mean + g.y_sd.fillna(0)).to_numpy(dtype=float)
            ax.fill_between(g.date, lo, hi, alpha=0.15, label="±1 SD" if (i == 0 and j == 0) else None)
            ax.set_title(f"{ds} — {model}")
            ax.grid(True, alpha=0.2)
            if j == 0:
                ax.set_ylabel("Observed / predicted value")
            ax.tick_params(axis="x", rotation=30, labelsize=7)
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.suptitle(f"Forecast visualization at H={horizon}, pre-specified central rolling origin ({origin_id})", y=0.995)
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(0.5, 0.955), ncol=3, frameon=False)
    fig.tight_layout(rect=[0, 0, 1, 0.91])
    fig.savefig(png, dpi=300, bbox_inches="tight")
    fig.savefig(pdf, bbox_inches="tight")
    plt.close(fig)


def make_figS1_originwise_mae(metrics: pd.DataFrame, png: Path, pdf: Path):
    fig, axes = plt.subplots(2, 2, figsize=(10.5, 7.2))
    horizons = sorted(metrics.horizon.unique())
    for ax, H in zip(axes.ravel(), horizons):
        for ds in ["Chennai", "Melbourne"]:
            # Unit scales differ, so normalize each model's origin MAE by the dataset/horizon Transformer MAE at the same origin.
            sub = metrics[(metrics.dataset == ds) & (metrics.horizon == H)]
            o = sub.groupby(["origin_id", "model"], as_index=False).agg(mae=("mae", "mean"))
            base = o[o.model == "Transformer"][["origin_id", "mae"]].rename(columns={"mae": "base"})
            for model in ["BLS", "A-BLAST"]:
                mm = o[o.model == model].merge(base, on="origin_id")
                rel = mm.mae / mm.base
                ax.plot(mm.origin_id, rel, marker="o", label=f"{ds} {model}/Transformer")
        ax.axhline(1.0, linestyle="--", linewidth=0.8)
        ax.set_title(f"H={H}")
        ax.set_ylabel("Relative MAE (Transformer = 1)")
        ax.grid(True, alpha=0.2)
    # Hide any unused panels in quick/smoke configurations with fewer than four horizons.
    for ax in axes.ravel()[len(horizons):]:
        ax.set_visible(False)
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.suptitle("Rolling-origin relative MAE stability", y=0.995)
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(0.5, 0.955), ncol=2, frameon=False)
    fig.tight_layout(rect=[0, 0, 1, 0.89])
    fig.savefig(png, dpi=300, bbox_inches="tight")
    fig.savefig(pdf, bbox_inches="tight")
    plt.close(fig)


def combine_tables_docx(paths, out_path: Path):
    doc = Document()
    _set_doc_font(doc)
    for idx, pth in enumerate(paths):
        src = Document(pth)
        if idx > 0:
            doc.add_page_break()
        for para in src.paragraphs:
            p = doc.add_paragraph(para.text)
            if para.runs and para.runs[0].bold:
                p.runs[0].bold = True
        for t in src.tables:
            nt = doc.add_table(rows=0, cols=len(t.columns))
            nt.style = "Table Grid"
            for row in t.rows:
                nr = nt.add_row().cells
                for i, c in enumerate(row.cells):
                    nr[i].text = c.text
    doc.save(out_path)
